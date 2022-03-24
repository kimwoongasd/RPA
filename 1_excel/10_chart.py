from openpyxl.chart.bar_chart import BarChart
from openpyxl.chart.reference import Reference
from openpyxl.chart.line_chart import LineChart
from openpyxl import load_workbook

wb = load_workbook("sample2.xlsx")
ws = wb.active

# b2 부터 c11 까지의 데이터를 차트로 생성
# bar_value = Reference(ws, min_row=2, max_row=11, min_col=2, max_col=3)
# bar_chart = BarChart() # 차트 종류 설정 (bar, line, pie...)
# bar_chart.add_data(bar_value) # 차트 데이터 추가

# # 차트 넣을 위치
# ws.add_chart(bar_chart, "E1")
line_value = Reference(ws, min_row=1, max_row=11, min_col=2, max_col=3)
line_chart = LineChart()
# titles_from_data=True : 계열 > 수학 영어 (제목에서 값을 가져옴)
line_chart.add_data(line_value, titles_from_data=True) 

# 제목
line_chart.title = '성적표'
# 미리 정의된 스타일 적용, 사용자가 개별 지정 가능
line_chart.style = 20
# y축 제목
line_chart.y_axis.title ="점수"
# x축 제목
line_chart.x_axis.title ='번호' 

ws.add_chart(line_chart, "E1")

wb.save("sample_chart.xlsx")
wb.close()